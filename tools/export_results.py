"""Export aggregate manuscript tables to portable UTF-8 CSV files."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
import shutil

from openpyxl import load_workbook


TABLE_SHEETS = {
    "T1 AFx4 fastMRI": "table_1_fastmri_r4.csv",
    "T2 AFx8 fastMRI": "table_2_fastmri_r8.csv",
    "T3 AFx12 fastMRI": "table_3_fastmri_r12.csv",
    "T4 Radial Spiral": "table_4_sampling_patterns.csv",
    "T5 Module Ablation": "table_5_architecture_ablation.csv",
    "T6 Config Ablation": "table_6_configuration_ablation.csv",
}
READER_SHEETS = {
    "Summary": "reader_summary.csv",
    "Table VII": "table_7_reader_assessment.csv",
    "Statistics": "reader_statistics.csv",
}


def write_sheet(sheet, path: Path) -> None:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            while row and row[-1] is None:
                row.pop()
            writer.writerow(["" if value is None else value for value in row])


def export_workbook(source: Path, sheets: dict[str, str], out: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    for sheet_name, file_name in sheets.items():
        write_sheet(workbook[sheet_name], out / file_name)
    workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tables", type=Path, required=True)
    parser.add_argument("--reliability", type=Path, required=True)
    parser.add_argument("--reader", type=Path, required=True)
    parser.add_argument("--prototype", type=Path, required=True)
    parser.add_argument("--out", type=Path, default=Path("results"))
    args = parser.parse_args()
    export_workbook(args.tables, TABLE_SHEETS, args.out)
    export_workbook(args.reliability, {"Results": "reliability_results.csv"}, args.out)
    export_workbook(args.reader, READER_SHEETS, args.out)
    args.out.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(args.prototype, args.out / "prototype_utilization.csv")
    print(f"Exported aggregate results to {args.out}")


if __name__ == "__main__":
    main()
